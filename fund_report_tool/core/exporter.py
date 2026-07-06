import os
import shutil
from pathlib import Path
from typing import Optional
import pandas as pd
from openpyxl import load_workbook

from core.database import get_session, ProjectFinancial
from core.template_matcher import strip_prefixes, _get_template_path
from core.paths import OUTPUT_DIR

# Extra aliases for DB item_name -> template row name
_EXPORT_ALIASES = {
    '总资产': '资产总计',
    '应收账款': '应收票据及应收账款',
    '应收票据': '应收票据及应收账款',
    '应收票据及应收账款': '应收票据及应收账款',
    '应付票据': '应付账款',
    '经营活动现金流量净额': '经营活动产生的现金流量净额',
    '投资活动现金流量净额': '投资活动产生的现金流量净额',
    '筹资活动现金流量净额': '筹资活动产生的现金流量净额',
    '现金及现金等价物净增加额': '五、现金及现金等价物净增加额',
    '期末现金及现金等价物余额': '六、期末现金及现金等价物余额',
}


def _normalize(text: str) -> str:
    """Normalize text for fuzzy matching."""
    if not text:
        return ''
    text = strip_prefixes(text)
    # Remove common suffixes like （大基金...）
    text = text.replace('（大基金）', '').replace('(大基金)', '')
    text = text.replace('（大基金部分为利息收入）', '')
    text = text.replace('/其中：应付股利', '')
    return text.strip()


def _build_template_index() -> dict:
    """Build a mapping from normalized template name -> (row_index, original_name)."""
    path = _get_template_path()
    if not path:
        return {}
    df = pd.read_excel(path, sheet_name=0, header=None)
    index = {}
    for i in range(len(df)):
        name = str(df.iloc[i, 0]).strip() if pd.notna(df.iloc[i, 0]) else ''
        if name and name not in ('项目名称', '所属时间', 'nan', 'None', ''):
            norm = _normalize(name)
            index[norm] = (i, name)
            index[name] = (i, name)  # also keep original
    return index


def _find_row_index(item_name: str, template_index: dict) -> Optional[int]:
    """Find template row index for a DB item_name."""
    if not item_name:
        return None

    # Direct alias
    if item_name in _EXPORT_ALIASES:
        alias_target = _EXPORT_ALIASES[item_name]
        if alias_target in template_index:
            return template_index[alias_target][0]

    # Exact match
    if item_name in template_index:
        return template_index[item_name][0]

    # Normalized match
    norm = _normalize(item_name)
    if norm in template_index:
        return template_index[norm][0]

    # Try alias normalized
    if item_name in _EXPORT_ALIASES:
        alias_target = _EXPORT_ALIASES[item_name]
        norm_alias = _normalize(alias_target)
        if norm_alias in template_index:
            return template_index[norm_alias][0]

    return None


def export_project_report(project_name: str, period: str, output_dir = OUTPUT_DIR) -> str:
    """Export project financial data into a template-formatted xlsx file."""
    template_path = _get_template_path()
    if not template_path:
        raise FileNotFoundError('Template file not found')

    session = get_session()
    try:
        items = session.query(ProjectFinancial).filter_by(
            project_name=project_name, period=period
        ).all()
    finally:
        session.close()

    if not items:
        raise ValueError(f'No data found for {project_name} {period}')

    # Build template index
    template_index = _build_template_index()

    # Copy template
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_project = project_name.replace('/', '_').replace('\\', '_')
    output_path = output_dir / f'{safe_project}_{period}_财报.xlsx'
    shutil.copy(template_path, str(output_path))

    # Load and fill
    wb = load_workbook(output_path)
    ws = wb.active

    # Fill period
    ws.cell(row=2, column=2, value=period)

    # Fill values
    for item in items:
        row_idx = _find_row_index(item.item_name, template_index)
        if row_idx is not None:
            # openpyxl uses 1-based row numbers; DataFrame was 0-based, so add 1
            excel_row = row_idx + 1
            ws.cell(row=excel_row, column=2, value=item.value)

    wb.save(str(output_path))
    return str(output_path)
