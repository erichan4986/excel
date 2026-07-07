import os
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from datetime import datetime
from core.database import get_session, ProjectFinancial, FundFinancial
from core import matcher
from core.paths import OUTPUT_DIR
from core.config import load_config


def _get_writable_cell(ws, row, col):
    """Return a writable cell, skipping MergedCell objects.
    If the target cell is part of a merged range, return the top-left
    cell of that range so we can write into it."""
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    # Find the merged range that contains this cell
    for merged_range in ws.merged_cells.ranges:
        if (merged_range.min_col <= col <= merged_range.max_col and
                merged_range.min_row <= row <= merged_range.max_row):
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return None


_NON_COMPANY_KEYWORDS = {
    "万元", "元", "千元", "百万元", "亿元",
    "单位", "指标", "项目", "名称", "计量", "计量单位",
    "合计", "总计", "期末", "期初", "本期", "上期", "年初", "年末", "金额",
    "资产负债表", "利润表", "现金流量表", "报表", "附注", "注释",
    "序号", "编号", "备注", "说明", "比例", "%", "百分比",
    "一", "二", "三", "四", "五", "六", "七", "八", "九", "十",
    "十一", "十二", "十三", "十四", "十五",
    "被投企业", "资产名称", "企业名称", "公司名称", "项目名", "项目编号",
    "是否已完全退出", "资产类型", "购买成本", "风险暴露", "持股比例",
    "应收利息", "五级分类", "账龄", "所属行业", "所属子行业",
    "投资日期", "投资成本", "已实现价值", "未实现价值", "组织形式",
    "项目注册地", "境内企业填写省份；境外企业填写国家或地区",
}


def _is_likely_company_name(name):
    """Heuristic filter to exclude common header labels from being treated as company names.
    Allows single-character codes (A, B, C) which are often used as project/company identifiers."""
    if not name or len(name) == 0:
        return False
    if name.replace('.', '').replace('-', '').replace(',', '').replace('，', '').isdigit():
        return False
    if name in _NON_COMPANY_KEYWORDS:
        return False
    # Exclude generic labels like "第X列", "第X行", "第X章"
    if name.startswith("第") and len(name) <= 4 and name[-1] in ("列", "行", "章", "节"):
        return False
    return True


def detect_period_in_template(ws):
    """Try to detect target period from template headers."""
    for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
        for cell in row:
            if cell and isinstance(cell, str):
                patterns = [
                    (r'(20\d{2})[年/-](\d{1,2})[月/-](\d{1,2})', lambda m: f"{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}"),
                    (r'(20\d{2})[年/-]?[Qq]([1-4])', lambda m: f"{m.group(1)}-{['03-31','06-30','09-30','12-31'][int(m.group(2))-1]}"),
                    (r'(20\d{2})[年/-](\d{1,2})', lambda m: f"{m.group(1)}-{m.group(2).zfill(2)}-30"),
                    (r'(20\d{2})', lambda m: f"{m.group(1)}-12-31"),
                ]
                for pattern, formatter in patterns:
                    match = re.search(pattern, cell)
                    if match:
                        return formatter(match)
    return None


def get_latest_period(db_path=None):
    """Get the latest period from database."""
    session = get_session(db_path=db_path)
    try:
        result = session.query(ProjectFinancial.period).distinct().order_by(ProjectFinancial.period.desc()).first()
        if result:
            return result[0]
        result = session.query(FundFinancial.period).distinct().order_by(FundFinancial.period.desc()).first()
        if result:
            return result[0]
        return None
    finally:
        session.close()


def get_fund_projects():
    return load_config().get('fund_projects', {})


def fill_template(template_path, mapping_overrides=None, db_path=None):
    """兼容旧接口，自动检测布局。"""
    output_path, review_path, _ = fill_template_with_layout(
        template_path,
        layout={},
        mapping_overrides=mapping_overrides,
        db_path=db_path
    )
    return output_path, review_path


def fill_template_with_layout(template_path, layout=None, period=None, mapping_overrides=None, db_path=None, use_llm=True):
    """
    支持显式二维布局的填表函数。

    layout 格式:
    {
        "company_axis": "row"|"col",
        "metric_axis": "row"|"col",
        "company_header_row": int,
        "company_header_col": int,
        "metric_header_row": int,
        "metric_header_col": int,
    }
    如果 layout 为空，则回退到原来的单维度列头扫描逻辑。
    """
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(template_path)
    ws = wb.active

    if not period:
        period = detect_period_in_template(ws)
    if not period:
        period = get_latest_period(db_path=db_path)
    if not period:
        period = "unknown"

    if not layout:
        return _fill_legacy(ws, wb, template_path, mapping_overrides, period, db_path, use_llm=use_llm)

    # Guard against None values from LLM-parsed layouts
    layout["company_header_row"] = layout.get("company_header_row") or 1
    layout["company_header_col"] = layout.get("company_header_col") or 1
    layout["metric_header_row"] = layout.get("metric_header_row") or 1
    layout["metric_header_col"] = layout.get("metric_header_col") or 1

    company_axis = layout.get("company_axis")
    metric_axis = layout.get("metric_axis")

    # Clear LLM match log before scanning so we only record matches for this fill
    matcher.clear_llm_match_log()

    def _best_match_row(ws, min_row, max_row):
        """Find the row with the most recognized metric headers."""
        best_row = min_row
        best_count = 0
        for row in range(min_row, max(min_row, max_row + 1)):
            count = 0
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value and isinstance(cell.value, str):
                    text = cell.value.strip()
                    if text and matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides):
                        count += 1
            if count > best_count:
                best_count = count
                best_row = row
        return best_row, best_count

    def _best_match_col(ws, min_col, max_col):
        """Find the column with the most recognized metric headers."""
        best_col = min_col
        best_count = 0
        for col in range(min_col, max(min_col, max_col + 1)):
            count = 0
            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value and isinstance(cell.value, str):
                    text = cell.value.strip()
                    if text and matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides):
                        count += 1
            if count > best_count:
                best_count = count
                best_col = col
        return best_col, best_count

    def _best_company_col(ws, min_col, max_col):
        """Find the column with the most company-like names."""
        best_col = min_col
        best_count = 0
        for col in range(min_col, max(min_col, max_col + 1)):
            count = 0
            for row in range(1, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value and isinstance(cell.value, str):
                    name = cell.value.strip()
                    if _is_likely_company_name(name):
                        count += 1
            if count > best_count:
                best_count = count
                best_col = col
        return best_col, best_count

    def _best_company_row(ws, min_row, max_row):
        """Find the row with the most company-like names."""
        best_row = min_row
        best_count = 0
        for row in range(min_row, max(min_row, max_row + 1)):
            count = 0
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=col)
                if cell.value and isinstance(cell.value, str):
                    name = cell.value.strip()
                    if _is_likely_company_name(name):
                        count += 1
            if count > best_count:
                best_count = count
                best_row = row
        return best_row, best_count

    companies = {}
    metrics_map = {}

    # Diagnostic logs to help users debug recognition issues
    company_scan_log = []
    metric_scan_log = []

    # Determine whether user explicitly specified header positions (don't override)
    user_specified_company_row = layout.get("_user_specified_company_row", False)
    user_specified_company_col = layout.get("_user_specified_company_col", False)
    user_specified_metric_row = layout.get("_user_specified_metric_row", False)
    user_specified_metric_col = layout.get("_user_specified_metric_col", False)

    # ---- Scan company headers ----
    if company_axis == "col":
        ch_row = layout.get("company_header_row", 1)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=ch_row, column=col)
            if cell.value and isinstance(cell.value, str):
                name = cell.value.strip()
                is_company = _is_likely_company_name(name)
                company_scan_log.append({
                    "pos": f"R{ch_row}C{col}", "raw": name,
                    "recognized": is_company,
                    "reason": "filtered" if not is_company else "ok"
                })
                if is_company:
                    companies[(ch_row, col)] = name
        # Fallback: if no companies found, auto-search nearby rows
        if len(companies) == 0:
            best_row, _ = _best_company_row(ws, max(1, ch_row - 2), ch_row + 2)
            if best_row != ch_row:
                companies = {}
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=best_row, column=col)
                    if cell.value and isinstance(cell.value, str):
                        name = cell.value.strip()
                        if _is_likely_company_name(name):
                            companies[(best_row, col)] = name
                layout["company_header_row"] = best_row
    elif company_axis == "row":
        ch_col = layout.get("company_header_col", 1)
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=ch_col)
            if cell.value and isinstance(cell.value, str):
                name = cell.value.strip()
                is_company = _is_likely_company_name(name)
                company_scan_log.append({
                    "pos": f"R{row}C{ch_col}", "raw": name,
                    "recognized": is_company,
                    "reason": "filtered" if not is_company else "ok"
                })
                if is_company:
                    companies[(row, ch_col)] = name
        # Fallback: if no companies found, auto-search nearby columns
        if len(companies) == 0 and not user_specified_company_col:
            best_col, _ = _best_company_col(ws, max(1, ch_col - 2), ch_col + 2)
            if best_col != ch_col:
                companies = {}
                for row in range(1, ws.max_row + 1):
                    cell = ws.cell(row=row, column=best_col)
                    if cell.value and isinstance(cell.value, str):
                        name = cell.value.strip()
                        if _is_likely_company_name(name):
                            companies[(row, best_col)] = name
                layout["company_header_col"] = best_col

    # ---- Scan metric headers ----
    if metric_axis == "col":
        mh_row = layout.get("metric_header_row", 1)
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=mh_row, column=col)
            if cell.value and isinstance(cell.value, str):
                text = cell.value.strip()
                if text:
                    std = matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides)
                    matched = bool(std) or (mapping_overrides and text in mapping_overrides)
                    metric_scan_log.append({
                        "pos": f"R{mh_row}C{col}", "raw": text,
                        "matched": matched,
                        "standard": std or (mapping_overrides.get(text) if mapping_overrides else None)
                    })
                    if std:
                        metrics_map[(mh_row, col)] = std
                    elif mapping_overrides and text in mapping_overrides:
                        metrics_map[(mh_row, col)] = mapping_overrides[text]
        # Fallback: if no metrics matched, auto-search nearby rows
        if len(metrics_map) == 0:
            best_row, best_count = _best_match_row(ws, max(1, mh_row - 2), mh_row + 2)
            if best_count > 0 and best_row != mh_row:
                metrics_map = {}
                for col in range(1, ws.max_column + 1):
                    cell = ws.cell(row=best_row, column=col)
                    if cell.value and isinstance(cell.value, str):
                        text = cell.value.strip()
                        if text:
                            std = matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides)
                            if std:
                                metrics_map[(best_row, col)] = std
                            elif mapping_overrides and text in mapping_overrides:
                                metrics_map[(best_row, col)] = mapping_overrides[text]
                layout["metric_header_row"] = best_row
    elif metric_axis == "row":
        mh_col = layout.get("metric_header_col", 1)
        for row in range(1, ws.max_row + 1):
            cell = ws.cell(row=row, column=mh_col)
            if cell.value and isinstance(cell.value, str):
                text = cell.value.strip()
                if text:
                    std = matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides)
                    matched = bool(std) or (mapping_overrides and text in mapping_overrides)
                    metric_scan_log.append({
                        "pos": f"R{row}C{mh_col}", "raw": text,
                        "matched": matched,
                        "standard": std or (mapping_overrides.get(text) if mapping_overrides else None)
                    })
                    if std:
                        metrics_map[(row, mh_col)] = std
                    elif mapping_overrides and text in mapping_overrides:
                        metrics_map[(row, mh_col)] = mapping_overrides[text]
        # Fallback: if no metrics matched, auto-search nearby columns
        if len(metrics_map) == 0:
            best_col, best_count = _best_match_col(ws, max(1, mh_col - 2), mh_col + 2)
            if best_count > 0 and best_col != mh_col:
                metrics_map = {}
                for row in range(1, ws.max_row + 1):
                    cell = ws.cell(row=row, column=best_col)
                    if cell.value and isinstance(cell.value, str):
                        text = cell.value.strip()
                        if text:
                            std = matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides)
                            if std:
                                metrics_map[(row, best_col)] = std
                            elif mapping_overrides and text in mapping_overrides:
                                metrics_map[(row, best_col)] = mapping_overrides[text]
                layout["metric_header_col"] = best_col

    stats = {
        "companies_found": len(companies),
        "metrics_found": len([v for v in metrics_map.values() if v is not None]),
        "filled_count": 0,
        "skipped_no_data": 0,
        "company_names": sorted(set(companies.values())),
        "metric_names": sorted(v for v in set(metrics_map.values()) if v is not None),
        "layout_used": {
            "company_axis": company_axis,
            "metric_axis": metric_axis,
            "company_header_row": layout.get("company_header_row"),
            "company_header_col": layout.get("company_header_col"),
            "metric_header_row": layout.get("metric_header_row"),
            "metric_header_col": layout.get("metric_header_col"),
        },
        "period_used": period,
        "company_scan_log": company_scan_log,
        "metric_scan_log": metric_scan_log,
        "llm_matches": matcher.get_llm_match_log(),
    }

    session = get_session(db_path=db_path)
    try:
        # Build a fuzzy company name lookup cache
        # e.g. template "A" → DB "公司A", "B" → "公司B"
        all_db_companies = [r[0] for r in session.query(ProjectFinancial.project_name).distinct().all()]

        def _resolve_company_name(template_name):
            if not template_name:
                return None
            # Exact match
            if template_name in all_db_companies:
                return template_name
            # Substring match: template name inside DB name ("A" in "公司A")
            for db_name in all_db_companies:
                if template_name in db_name:
                    return db_name
            # Substring match: DB name inside template name ("公司A" in "公司A项目")
            for db_name in all_db_companies:
                if db_name in template_name:
                    return db_name
            return template_name

        company_name_cache = {}

        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                if (r, c) in companies or (r, c) in metrics_map:
                    continue

                company_name = None
                metric_name = None

                if company_axis == "col":
                    company_name = companies.get((layout.get("company_header_row", 1), c))
                elif company_axis == "row":
                    company_name = companies.get((r, layout.get("company_header_col", 1)))

                # In 2D layout mode, skip cells that don't have a corresponding company.
                # Prevents filling empty rows (e.g. row 6/7 with no company name)
                # with unfiltered data from arbitrary projects.
                if company_axis and not company_name:
                    continue

                if metric_axis == "col":
                    metric_name = metrics_map.get((layout.get("metric_header_row", 1), c))
                elif metric_axis == "row":
                    metric_name = metrics_map.get((r, layout.get("metric_header_col", 1)))

                if not metric_name:
                    continue

                # Resolve fuzzy company name match
                if company_name and company_name not in company_name_cache:
                    company_name_cache[company_name] = _resolve_company_name(company_name)
                resolved_name = company_name_cache.get(company_name) if company_name else None

                value = None
                source_note = ""
                if resolved_name:
                    result = session.query(ProjectFinancial).filter_by(
                        project_name=resolved_name, period=period, item_name=metric_name
                    ).first()
                    if result:
                        value = result.value
                        source_note = f"项目:{resolved_name}, 期间:{period}"
                else:
                    result = session.query(ProjectFinancial).filter_by(
                        period=period, item_name=metric_name
                    ).first()
                    if result:
                        value = result.value
                        source_note = f"项目:{result.project_name}, 期间:{period}"

                if value is not None:
                    cell = _get_writable_cell(ws, r, c)
                    if cell is None:
                        continue
                    cell.value = value
                    comment = Comment(
                        f"数据来源: {source_note}\n填充时间: {datetime.now().isoformat()}",
                        "FundReportTool"
                    )
                    cell.comment = comment
                    stats["filled_count"] += 1
                else:
                    stats["skipped_no_data"] += 1

        base_name = os.path.splitext(os.path.basename(template_path))[0]
        output_path = OUTPUT_DIR / f"{base_name}_filled.xlsx"
        review_path = OUTPUT_DIR / f"{base_name}_review.xlsx"
        wb.save(str(output_path))
        wb.save(str(review_path))
        return str(output_path), str(review_path), stats
    finally:
        session.close()


def _fill_legacy(ws, wb, template_path, mapping_overrides, period, db_path, use_llm=True):
    """保留原来的单维度列头填充逻辑。"""
    headers = {}
    header_row = 1
    for col_idx, cell in enumerate(ws[header_row], 1):
        if cell.value and isinstance(cell.value, str):
            text = cell.value.strip()
            if len(text) < 30 and not text.replace('.', '').replace('-', '').isdigit():
                headers[(header_row, col_idx)] = text

    if not headers:
        for row_idx in range(1, min(6, ws.max_row + 1)):
            for col_idx, cell in enumerate(ws[row_idx], 1):
                if cell.value and isinstance(cell.value, str):
                    text = cell.value.strip()
                    if len(text) < 30 and not text.replace('.', '').replace('-', '').isdigit():
                        headers[(row_idx, col_idx)] = text
                        header_row = row_idx

    matched = {}
    for (row, col), text in headers.items():
        standard = matcher.match_header(text, use_llm=use_llm, mapping_overrides=mapping_overrides)
        if standard:
            matched[(row, col)] = standard

    if mapping_overrides:
        for key, val in mapping_overrides.items():
            for (row, col), text in headers.items():
                if text == key and val:
                    matched[(row, col)] = val

    data_cells = {}
    for (row, col), standard in matched.items():
        for r in range(row + 1, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            if cell.value is not None:
                val = str(cell.value).strip()
                if val and not val.isalpha():
                    data_cells[(r, col)] = standard
                    break

    session = get_session(db_path=db_path)
    config = load_config()
    fund_projects = get_fund_projects()

    try:
        for (row, col), standard in data_cells.items():
            mapping_config = None
            for m in config.get('mappings', []):
                if m['standard'] == standard:
                    mapping_config = m
                    break

            value = None
            source_note = ""

            if mapping_config and mapping_config.get('aggregation') == 'sum':
                fund_name = list(fund_projects.keys())[0] if fund_projects else None
                projects = fund_projects.get(fund_name, [])
                total = 0.0
                for proj in projects:
                    result = session.query(ProjectFinancial).filter_by(
                        project_name=proj, period=period, item_name=standard
                    ).first()
                    if result and result.value is not None:
                        total += result.value
                value = total if total != 0 else None
                source_note = f"基金聚合({fund_name}): {len(projects)}个项目求和, 期间:{period}"
            else:
                result = session.query(ProjectFinancial).filter_by(
                    period=period, item_name=standard
                ).first()
                if result:
                    value = result.value
                    source_note = f"项目数据:{result.project_name}, 期间:{period}"
                else:
                    result = session.query(FundFinancial).filter_by(
                        period=period, item_name=standard
                    ).first()
                    if result:
                        value = result.value
                        source_note = f"基金数据, 期间:{period}"

            if value is not None:
                cell = _get_writable_cell(ws, row, col)
                if cell is None:
                    continue
                cell.value = value
                comment = Comment(
                    f"数据来源: {source_note}\n填充时间: {datetime.now().isoformat()}",
                    "FundReportTool"
                )
                cell.comment = comment

        base_name = os.path.splitext(os.path.basename(template_path))[0]
        output_path = OUTPUT_DIR / f"{base_name}_filled.xlsx"
        review_path = OUTPUT_DIR / f"{base_name}_review.xlsx"
        wb.save(str(output_path))
        wb.save(str(review_path))
        stats = {
            "companies_found": 0,
            "metrics_found": len(matched),
            "filled_count": sum(1 for r in data_cells if ws.cell(row=r[0], column=r[1]).value is not None),
            "skipped_no_data": 0,
            "company_names": [],
            "metric_names": sorted(set(matched.values())),
        }
        return str(output_path), str(review_path), stats
    finally:
        session.close()
