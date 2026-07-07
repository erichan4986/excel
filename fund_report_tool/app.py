import os
from pathlib import Path
from typing import List, Optional
from datetime import datetime
from uuid import uuid4
from dotenv import load_dotenv
load_dotenv()
import shutil
import json
import yaml
from fastapi import FastAPI, File, UploadFile, Form, HTTPException, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from sqlalchemy import func

from core.database import init_db, get_session, ProjectFinancial, ProjectMetric, FundFairValue, UploadRecord
from core import cleaner
from core.cleaner import clean_and_store
from core.filler import fill_template
from core.matcher import match_header, batch_match, get_mappings
from core.config import load_config
from core.llm_helper import generate_mapping_suggestion
from core.pdf_processor import extract_tables
from core import chatbot
from core.exporter import export_project_report
from core.fund_exporter import export_fund_report
from core.paths import (
    UPLOAD_DIR, CLEANED_DIR, OUTPUT_DIR, STATIC_DIR, TEMPLATES_DIR, CONFIG_PATH
)


def _sanitize_filename(filename: Optional[str]) -> str:
    """Strip directory components from uploaded filenames."""
    if not filename:
        return "unnamed"
    name = Path(filename).name
    if not name or name in (".", ".."):
        return "unnamed"
    return name


def _is_safe_path(base_dir: Path, target_path: Path) -> bool:
    """Check that target_path is inside base_dir (after resolving symlinks)."""
    try:
        target_path.resolve().relative_to(base_dir.resolve())
        return True
    except ValueError:
        return False


def _save_upload(file: UploadFile) -> tuple[str, str, Path]:
    """Save an uploaded file with a unique stored name. Returns (original_name, stored_name, file_path)."""
    original_name = _sanitize_filename(file.filename)
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    unique = str(uuid4())[:8]
    stored_name = f"{timestamp}_{unique}_{original_name}"
    file_path = UPLOAD_DIR / stored_name
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    return original_name, stored_name, file_path


def _record_upload(original_name: str, stored_name: str, stored_path: Path):
    """Record upload metadata in the database."""
    session = get_session()
    try:
        session.add(UploadRecord(
            original_name=original_name,
            stored_name=stored_name,
            stored_path=str(stored_path)
        ))
        session.commit()
    finally:
        session.close()


app = FastAPI(title="基金投后数据自动清洗填表系统")

for d in [UPLOAD_DIR, CLEANED_DIR, OUTPUT_DIR, STATIC_DIR, TEMPLATES_DIR]:
    d.mkdir(parents=True, exist_ok=True)

app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
templates = Jinja2Templates(directory=str(TEMPLATES_DIR))


@app.on_event("startup")
async def startup():
    init_db()


@app.get("/")
async def index(request: Request):
    return templates.TemplateResponse(request, "index.html", {})


@app.post("/api/clean")
async def api_clean(file: UploadFile = File(...), data_type: str = Form(...), overwrite: bool = Form(False), period_override: Optional[str] = Form(None)):
    if data_type not in ('project', 'fund'):
        return JSONResponse({"error": "Invalid data_type, must be 'project' or 'fund'"}, status_code=400)

    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    result = clean_and_store(
        str(file_path), data_type, overwrite=overwrite,
        original_filename=original_name, period_override=period_override
    )
    result['original_name'] = original_name
    result['stored_name'] = stored_name
    result['stored_path'] = str(file_path)
    return JSONResponse(result)


@app.post("/api/clean-batch")
async def api_clean_batch(files: List[UploadFile] = File(...), data_type: str = Form(...), overwrite: bool = Form(False), period_override: Optional[str] = Form(None)):
    if data_type not in ('project', 'fund'):
        return JSONResponse({"error": "Invalid data_type, must be 'project' or 'fund'"}, status_code=400)

    results = []
    for file in files:
        original_name, stored_name, file_path = _save_upload(file)
        _record_upload(original_name, stored_name, file_path)
        result = clean_and_store(
            str(file_path), data_type, overwrite=overwrite,
            original_filename=original_name, period_override=period_override
        )
        result["filename"] = original_name
        result["stored_name"] = stored_name
        result["stored_path"] = str(file_path)
        results.append(result)

    success_count = sum(1 for r in results if r["status"] == "success")
    warning_count = sum(1 for r in results if r["status"] == "warning")
    error_count = sum(1 for r in results if r["status"] == "error")
    duplicates = [r for r in results if r["status"] == "confirm"]

    if len(duplicates) == len(results) and len(results) > 0:
        return JSONResponse({
            "total": len(results),
            "success": 0,
            "warning": 0,
            "error": 0,
            "duplicates": len(duplicates),
            "status": "confirm",
            "results": results
        })

    return JSONResponse({
        "total": len(results),
        "success": success_count,
        "warning": warning_count,
        "error": error_count,
        "duplicates": len(duplicates),
        "results": results
    })


@app.post("/api/parse_template")
async def api_parse_template(file: UploadFile = File(...)):
    from openpyxl import load_workbook
    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    try:
        wb = load_workbook(str(file_path), data_only=True)
        ws = wb.active

        headers = []
        for row_idx in range(1, min(6, ws.max_row + 1)):
            for col_idx, cell in enumerate(ws[row_idx], 1):
                if cell.value and isinstance(cell.value, str):
                    text = cell.value.strip()
                    if len(text) < 30 and not text.replace('.', '').replace('-', '').isdigit():
                        headers.append(text)

        matched = {}
        missing = []
        for h in headers:
            std = match_header(h, use_llm=False)
            if std:
                matched[h] = std
            else:
                missing.append(h)

        return JSONResponse({
            "headers": headers,
            "matched": matched,
            "missing": missing,
            "original_name": original_name,
            "stored_name": stored_name,
            "stored_path": str(file_path)
        })
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/fill")
async def api_fill(file: UploadFile = File(...), mapping_overrides: str = Form("{}")):
    overrides = json.loads(mapping_overrides) if mapping_overrides else {}
    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    try:
        output_path, review_path = fill_template(str(file_path), overrides)
        return JSONResponse({
            "output": os.path.basename(output_path),
            "review": os.path.basename(review_path),
            "original_name": original_name,
            "stored_name": stored_name,
            "stored_path": str(file_path)
        })
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.get("/download/{filename}")
async def download(filename: str):
    safe_name = _sanitize_filename(filename)
    file_path = (OUTPUT_DIR / safe_name).resolve()
    if not _is_safe_path(OUTPUT_DIR, file_path):
        raise HTTPException(status_code=403, detail="Access denied")
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    return FileResponse(str(file_path), filename=safe_name)


@app.post("/api/smart-match")
async def api_smart_match(headers: list[str]):
    result = batch_match(headers)
    return JSONResponse(result)


@app.get("/api/mappings")
async def api_get_mappings():
    return JSONResponse({"mappings": get_mappings()})


@app.post("/api/mappings")
async def api_update_mappings(mappings: list[dict]):
    config = load_config()
    config['mappings'] = mappings
    with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
        yaml.dump(config, f, allow_unicode=True, sort_keys=False)
    return JSONResponse({"status": "success"})


@app.get("/api/projects")
async def api_projects():
    session = get_session()
    try:
        project_names = [
            r[0] for r in session.query(ProjectFinancial.project_name).distinct().all()
        ]
        return JSONResponse({"projects": project_names})
    finally:
        session.close()


@app.get("/api/periods")
async def api_periods(project: str):
    session = get_session()
    try:
        periods = [
            r[0] for r in session.query(ProjectFinancial.period)
            .filter_by(project_name=project).distinct().order_by(ProjectFinancial.period.desc()).all()
        ]
        return JSONResponse({"periods": periods})
    finally:
        session.close()


@app.get("/api/data")
async def api_data(project: str, period: str):
    session = get_session()
    try:
        items = session.query(ProjectFinancial).filter_by(
            project_name=project, period=period
        ).order_by(ProjectFinancial.id).all()
        return JSONResponse({
            "items": [
                {"item_name": i.item_name, "value": i.value, "report_type": i.report_type}
                for i in items
            ]
        })
    finally:
        session.close()


@app.get("/api/metrics_data")
async def api_metrics_data(project: str, period: str):
    session = get_session()
    try:
        metrics = session.query(ProjectMetric).filter_by(
            project_name=project, period=period
        ).order_by(ProjectMetric.id).all()
        return JSONResponse({
            "metrics": [
                {"metric_name": m.metric_name, "metric_value": m.metric_value}
                for m in metrics
            ]
        })
    finally:
        session.close()


@app.get("/api/download_report")
async def api_download_report(project: str, period: str):
    try:
        output_path = export_project_report(project, period)
        filename = os.path.basename(output_path)
        return FileResponse(output_path, filename=filename)
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/delete_project_period")
async def api_delete_project_period(request: Request):
    body = await request.json()
    project_name = body.get("project_name")
    period = body.get("period")

    if not project_name or not period:
        return JSONResponse({"error": "project_name and period are required"}, status_code=400)

    session = get_session()
    try:
        fin_deleted = session.query(ProjectFinancial).filter_by(
            project_name=project_name, period=period
        ).delete()
        metric_deleted = session.query(ProjectMetric).filter_by(
            project_name=project_name, period=period
        ).delete()
        session.commit()
        return JSONResponse({
            "status": "success",
            "financial_records_deleted": fin_deleted,
            "metric_records_deleted": metric_deleted
        })
    except Exception as e:
        session.rollback()
        return JSONResponse({"error": str(e)}, status_code=500)
    finally:
        session.close()


@app.get("/api/funds")
async def api_funds():
    session = get_session()
    try:
        fund_names = [
            r[0] for r in session.query(FundFairValue.fund_name).distinct().order_by(FundFairValue.fund_name).all()
        ]
        return JSONResponse({"funds": fund_names})
    finally:
        session.close()


@app.get("/api/fund_periods")
async def api_fund_periods(fund: str):
    session = get_session()
    try:
        periods = [
            r[0] for r in session.query(FundFairValue.period)
            .filter_by(fund_name=fund).distinct().order_by(FundFairValue.period.desc()).all()
        ]
        return JSONResponse({"periods": periods})
    finally:
        session.close()



@app.get("/api/fund_data")
async def api_fund_data(fund: str, period: str):
    session = get_session()
    try:
        items = session.query(FundFairValue).filter_by(
            fund_name=fund, period=period
        ).order_by(FundFairValue.id).all()

        total_cost = session.query(func.sum(FundFairValue.cost)).filter_by(
            fund_name=fund, period=period
        ).scalar() or 0.0

        total_fair_value = session.query(func.sum(FundFairValue.fair_value)).filter_by(
            fund_name=fund, period=period
        ).scalar() or 0.0

        return JSONResponse({
            "items": [
                {
                    "project_name": i.project_name,
                    "cost": i.cost,
                    "fair_value": i.fair_value,
                    "total_return": i.total_return,
                    "remark": i.remark,
                    "last_payment_date": i.last_payment_date,
                }
                for i in items
            ],
            "summary": {
                "total_cost": round(total_cost, 2),
                "total_fair_value": round(total_fair_value, 2),
            }
        })
    finally:
        session.close()


@app.post("/api/delete_fund_period")
async def api_delete_fund_period(request: Request):
    body = await request.json()
    fund_name = body.get("fund_name")
    period = body.get("period")

    if not fund_name or not period:
        return JSONResponse({"error": "fund_name and period are required"}, status_code=400)

    session = get_session()
    try:
        deleted = session.query(FundFairValue).filter_by(
            fund_name=fund_name, period=period
        ).delete()
        session.commit()
        return JSONResponse({
            "status": "success",
            "records_deleted": deleted
        })
    except Exception as e:
        session.rollback()
        return JSONResponse({"error": str(e)}, status_code=500)
    finally:
        session.close()


@app.get("/api/download_fund_report")
async def api_download_fund_report(fund: str, period: str):
    try:
        output_path = export_fund_report(fund, period)
        filename = os.path.basename(output_path)
        return FileResponse(output_path, filename=filename)
    except FileNotFoundError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/parse_pdf")
async def api_parse_pdf(file: UploadFile = File(...)):
    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    try:
        tables = extract_tables(str(file_path))
        return JSONResponse({
            "tables_count": len(tables),
            "tables": [t.head(10).to_dict(orient='records') for t in tables],
            "original_name": original_name,
            "stored_name": stored_name,
            "stored_path": str(file_path)
        })
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


@app.post("/api/delete_template")
async def api_delete_template(request: Request):
    body = await request.json()
    template_path = body.get("template_path", "")

    file_path = Path(template_path).resolve()
    if not _is_safe_path(UPLOAD_DIR, file_path):
        return JSONResponse({"error": "Access denied"}, status_code=403)
    if not file_path.exists():
        return JSONResponse({"error": "File not found"}, status_code=404)

    try:
        file_path.unlink()
        return JSONResponse({"status": "success"})
    except Exception as e:
        return JSONResponse({"error": str(e)}, status_code=500)


class ChatRequest(BaseModel):
    message: str
    context: dict = {}


class CleanConfirmRequest(BaseModel):
    file_path: str
    data_type: str
    confirmed_mappings: dict = {}
    overwrite: bool = False


@app.post("/api/chat")
async def api_chat(req: ChatRequest):
    parsed = chatbot.process_message(req.message, req.context)
    intent = parsed.get("intent", "general_chat")
    params = parsed.get("params", {})
    response_text = parsed.get("response", "")

    result_data = {}

    if intent == "fill_template":
        if not params.get("template_path") and req.context.get("template_path"):
            params["template_path"] = req.context["template_path"]
        exec_result = chatbot.execute_fill(params)
        result_data = exec_result
        if exec_result.get("success"):
            response_text += f" {exec_result.get('message', '')}"
        else:
            response_text += f" 填表失败：{exec_result.get('error')}"

    elif intent == "query_data":
        exec_result = chatbot.execute_query(params)
        result_data = exec_result
        if exec_result.get("success"):
            rows = exec_result.get("data", [])
            response_text += f" 查询到 {len(rows)} 条记录。"
        else:
            response_text += " 查询失败。"

    elif intent == "list_projects":
        exec_result = chatbot.execute_list_projects()
        result_data = exec_result
        response_text += f" 共 {len(exec_result.get('projects', []))} 个项目。"

    elif intent == "list_periods":
        exec_result = chatbot.execute_list_periods(params)
        result_data = exec_result
        response_text += f" 共 {len(exec_result.get('periods', []))} 个期间。"

    elif intent == "show_data":
        exec_result = chatbot.execute_show_data(params)
        result_data = exec_result
        response_text += f" 共 {len(exec_result.get('items', []))} 条财务记录。"

    elif intent == "show_fund_data":
        exec_result = chatbot.execute_fund_summary(params)
        result_data = exec_result
        if exec_result.get("success"):
            response_text += f" 该基金共投资 {exec_result.get('project_count', 0)} 个项目，总成本 {exec_result.get('total_cost', 0):,.2f}，总公允价值 {exec_result.get('total_fair_value', 0):,.2f}。"
        else:
            response_text += f" 查询失败：{exec_result.get('error')}"

    elif intent == "clarify_intent":
        result_data = {"options": params.get("options", [])}

    elif intent == "confirm_mappings":
        ctx = req.context or {}
        exec_result = chatbot.execute_confirm_mappings(ctx)
        result_data = exec_result
        if exec_result.get("success"):
            response_text = exec_result.get("message", "映射已保存。")
        else:
            response_text = f"保存失败：{exec_result.get('error')}"

    elif intent == "reject_mappings":
        ctx = req.context or {}
        exec_result = chatbot.execute_reject_mappings(ctx)
        result_data = exec_result
        if exec_result.get("success"):
            response_text = exec_result.get("message", "已跳过当前映射。")
        else:
            response_text = f"处理失败：{exec_result.get('error')}"

    elif intent == "draw_chart":
        exec_result = chatbot.execute_draw_chart(params)
        result_data = exec_result
        if exec_result.get("success"):
            rows = exec_result.get("series", [{}])[0].get("data", [])
            x_axis = exec_result.get("x_axis", [])
            response_text += f" 已绘制 {len(x_axis)} 个期间的 {exec_result.get('chart_type', 'line')} 图。"
        else:
            response_text += f" 绘图失败：{exec_result.get('error')}"

    elif intent == "calculate_metric":
        exec_result = chatbot.execute_calculate_metric(params)
        result_data = exec_result
        if exec_result.get("success"):
            response_text = exec_result.get("message", "计算完成。")
        else:
            response_text = f"计算失败：{exec_result.get('error')}"

    return JSONResponse({
        "intent": intent,
        "response": response_text,
        "data": result_data,
        "params": params
    })


@app.post("/api/clean-preview")
async def api_clean_preview(file: UploadFile = File(...), data_type: str = Form(...), period_override: Optional[str] = Form(None)):
    if data_type not in ('project', 'fund'):
        return JSONResponse({"error": "Invalid data_type, must be 'project' or 'fund'"}, status_code=400)

    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    result = cleaner.preview_clean(
        str(file_path), data_type,
        original_filename=original_name, period_override=period_override
    )
    result['original_name'] = original_name
    result['stored_name'] = stored_name
    result['stored_path'] = str(file_path)
    return JSONResponse(result)


@app.post("/api/clean-confirm")
async def api_clean_confirm(file: UploadFile = File(...), data_type: str = Form(...),
                            confirmed_mappings: str = Form("{}"), overwrite: bool = Form(False),
                            period_override: Optional[str] = Form(None)):
    if data_type not in ('project', 'fund'):
        return JSONResponse({"error": "Invalid data_type, must be 'project' or 'fund'"}, status_code=400)

    confirmed = json.loads(confirmed_mappings) if confirmed_mappings else {}

    original_name, stored_name, file_path = _save_upload(file)
    _record_upload(original_name, stored_name, file_path)

    # Persist confirmed mappings to config.yaml and skip_list
    config = load_config()

    existing_mappings = {m['standard']: m for m in config.get('mappings', [])}
    skip_list = set(config.get('skip_list', []))

    for raw, std in confirmed.items():
        if std is None:
            skip_list.add(raw)
        elif isinstance(std, str):
            if std in existing_mappings:
                aliases = set(existing_mappings[std].get('aliases', []))
                aliases.add(raw)
                existing_mappings[std]['aliases'] = sorted(aliases)
            else:
                existing_mappings[std] = {
                    "standard": std,
                    "aliases": [raw],
                    "category": "other",
                    "db_source": "project" if data_type == 'project' else "fund"
                }

    config['mappings'] = list(existing_mappings.values())
    config['skip_list'] = sorted(skip_list)
    with open(CONFIG_PATH, 'w', encoding='utf-8') as f:
        yaml.dump(config, f, allow_unicode=True, sort_keys=False)

    result = cleaner.clean_and_store(
        str(file_path), data_type, overwrite=overwrite, confirmed_mappings=confirmed,
        original_filename=original_name, period_override=period_override
    )
    result['original_name'] = original_name
    result['stored_name'] = stored_name
    result['stored_path'] = str(file_path)
    return JSONResponse(result)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)
