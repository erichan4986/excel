import os
import pandas as pd
from openpyxl import load_workbook
from core.database import init_db, get_session, ProjectFinancial
from core import filler

TEST_DB = "data/test_filler_debug.db"

def setup_module():
    for p in [TEST_DB]:
        if os.path.exists(p):
            os.remove(p)
    init_db(db_path=TEST_DB)
    session = get_session(db_path=TEST_DB)
    # Insert test data for 公司A 2025
    for item, val in [
        ("总资产", 1234.5),
        ("负债合计", 567.8),
        ("营业收入", 999.0),
    ]:
        session.add(ProjectFinancial(
            project_name="公司A", period="2025-12-31",
            report_type="BS", item_name=item, value=val,
            source_file="test.xlsx"
        ))
    session.commit()
    session.close()

def teardown_module():
    for p in [TEST_DB]:
        if os.path.exists(p):
            os.remove(p)

def test_fill_row_company_col_metric():
    """
    Template layout:
         A        B         C         D
    1   指标      单位     公司名    总资产
    2   总资产    万元     公司A     ?
    3   负债合计  万元     ?         ?
    
    company_axis='row' (company names down column C)
    metric_axis='col' (metrics across row 1)
    """
    df = pd.DataFrame({
        "指标": ["总资产", "负债合计"],
        "单位": ["万元", "万元"],
        "公司名": ["公司A", ""],
        "总资产": ["", ""],
    })
    path = "data/test_template_row_company.xlsx"
    df.to_excel(path, index=False)

    layout = {
        "company_axis": "row",
        "metric_axis": "col",
        "company_header_col": 3,
        "metric_header_row": 1,
    }
    out_path, review_path, stats = filler.fill_template_with_layout(
        path, layout=layout, period="2025-12-31", db_path=TEST_DB
    )
    print(f"Stats: {stats}")
    assert stats["metrics_found"] > 0, f"No metrics found: {stats}"
    assert stats["filled_count"] > 0, f"No cells filled: {stats}"

    wb = load_workbook(out_path)
    ws = wb.active
    # Cell D2 should be filled with 1234.5
    # openpyxl columns are 1-based: D=4, row=2
    val = ws.cell(row=2, column=4).value
    print(f"D2 value: {val}")
    assert val == 1234.5, f"Expected 1234.5, got {val}"


def test_fill_merged_cell_template():
    """
    Template with merged cells should not raise 'MergedCell' read-only error.
    Layout:
         A          B         C
    1   (empty)    公司A     公司B
    2   总资产     <merged>
    Where B2:C2 is merged. Writing to C2 should route to B2.
    """
    from openpyxl import Workbook
    path = "data/test_template_merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=2, value="公司A")
    ws.cell(row=1, column=3, value="公司B")
    ws.cell(row=2, column=1, value="总资产")
    ws.cell(row=2, column=2, value="")
    # Merge B2:C2 — C2 is a MergedCell, writing should route to B2
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=3)
    wb.save(path)

    layout = {
        "company_axis": "col",
        "metric_axis": "row",
        "company_header_row": 1,
        "metric_header_col": 1,
    }
    out_path, review_path, stats = filler.fill_template_with_layout(
        path, layout=layout, period="2025-12-31", db_path=TEST_DB
    )
    print(f"Merged cell stats: {stats}")
    assert stats["metrics_found"] > 0, f"No metrics found: {stats}"
    assert stats["filled_count"] > 0, f"No cells filled: {stats}"

    wb_out = load_workbook(out_path)
    ws_out = wb_out.active
    # B2 should contain 1234.5 (公司A's 总资产)
    val = ws_out.cell(row=2, column=2).value
    print(f"B2 value after merged fill: {val}")
    assert val == 1234.5, f"Expected 1234.5, got {val}"


def test_auto_detect_metric_row_when_wrong_row_given():
    """
    User says 'metric header in row 2' but actual metrics are in row 1.
    Row 2 has no recognizable metrics (only generic labels), so auto-detection
    should kick in and find row 1.
    Template:
         A          B         C         D
    1   指标        单位      总资产    负债合计
    2   项目名称    计量单位  公司A     ?
    3   负债合计    万元      ?         ?
    """
    from openpyxl import Workbook
    path = "data/test_template_auto_detect.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="指标")
    ws.cell(row=1, column=2, value="单位")
    ws.cell(row=1, column=3, value="总资产")
    ws.cell(row=1, column=4, value="负债合计")
    ws.cell(row=2, column=1, value="项目名称")
    ws.cell(row=2, column=2, value="计量单位")
    ws.cell(row=2, column=3, value="公司A")
    ws.cell(row=2, column=4, value="")
    ws.cell(row=3, column=1, value="负债合计")
    ws.cell(row=3, column=2, value="万元")
    ws.cell(row=3, column=3, value="")
    ws.cell(row=3, column=4, value="")
    wb.save(path)

    layout = {
        "company_axis": "row",
        "metric_axis": "col",
        "company_header_col": 3,
        "metric_header_row": 2,  # WRONG: actual metrics are in row 1
    }
    out_path, review_path, stats = filler.fill_template_with_layout(
        path, layout=layout, period="2025-12-31", db_path=TEST_DB
    )
    print(f"Auto-detect stats: {stats}")
    assert stats["metrics_found"] > 0, f"Auto-detect failed, no metrics: {stats}"

    wb_out = load_workbook(out_path)
    ws_out = wb_out.active
    # D2 should be filled because auto-detection found row 1 has "总资产" and "负债合计"
    val = ws_out.cell(row=2, column=4).value
    print(f"D2 value after auto-detect: {val}")
    # D2 gets 负债合计 for 公司A = 567.8
    assert val == 567.8, f"Expected 567.8, got {val}"
