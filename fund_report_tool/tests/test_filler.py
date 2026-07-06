import os
import pytest
from openpyxl import Workbook
from core.database import init_db, get_session, ProjectFinancial
from core import filler

TEST_DB = "data/test_filler.db"


@pytest.fixture(autouse=True)
def setup():
    for p in [TEST_DB]:
        if os.path.exists(p):
            os.remove(p)
    init_db(db_path=TEST_DB)
    os.makedirs("data/outputs", exist_ok=True)

    session = get_session(db_path=TEST_DB)
    session.add(ProjectFinancial(
        project_name="项目A", period="2024-12-31",
        report_type="资产负债表", item_name="总资产", value=999.0,
        source_file="test.xlsx"
    ))
    session.commit()
    session.close()
    yield
    for p in [TEST_DB]:
        if os.path.exists(p):
            os.remove(p)


def make_template(path):
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="指标")
    ws.cell(row=1, column=2, value="2024年12月")
    ws.cell(row=2, column=1, value="总资产")
    ws.cell(row=2, column=2, value=None)
    wb.save(path)


def test_fill_template():
    template = "data/test_template.xlsx"
    make_template(template)
    out, rev = filler.fill_template(template, db_path=TEST_DB)
    assert os.path.exists(out)
    assert os.path.exists(rev)


def test_fill_with_layout_company_col():
    session = get_session(db_path=TEST_DB)
    session.add(ProjectFinancial(
        project_name="公司A", period="2024-12-31",
        report_type="BS", item_name="总资产", value=1000.0,
        source_file="t.xlsx"
    ))
    session.add(ProjectFinancial(
        project_name="公司B", period="2024-12-31",
        report_type="BS", item_name="总资产", value=2000.0,
        source_file="t.xlsx"
    ))
    session.commit()
    session.close()

    template = "data/test_matrix.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, "项目/公司")
    ws.cell(1, 2, "公司A")
    ws.cell(1, 3, "公司B")
    ws.cell(2, 1, "总资产")
    ws.cell(2, 2, None)
    ws.cell(2, 3, None)
    wb.save(template)

    layout = {
        "company_axis": "col",
        "metric_axis": "row",
        "company_header_row": 1,
        "company_header_col": 2,
        "metric_header_row": 2,
        "metric_header_col": 1
    }
    out, rev, stats = filler.fill_template_with_layout(template, layout=layout, period="2024-12-31", db_path=TEST_DB)
    assert os.path.exists(out)
    assert stats["filled_count"] == 2

    from openpyxl import load_workbook
    wb_out = load_workbook(out)
    ws_out = wb_out.active
    assert ws_out.cell(2, 2).value == 1000.0
    assert ws_out.cell(2, 3).value == 2000.0
